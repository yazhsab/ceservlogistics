"""Build the CESERVE end-user Word guides from the adjacent Markdown sources.

Run with the Codex bundled Python runtime. PDF and visual review use the
Documents skill render_docx.py after this builder completes.
"""
from pathlib import Path
import re
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = ROOT / "output" / "docx"
FONT = "Liberation Sans"
GUIDES = [
    ("01-getting-started", "CESERVE_01_Getting_Started_and_Booking", "Getting started and booking"),
    ("02-administrator-configuration", "CESERVE_02_Administrator_Configuration", "Administrator configuration"),
    ("03-daily-operations", "CESERVE_03_Daily_Operations", "Daily operations"),
    ("04-common-questions", "CESERVE_04_Common_Questions", "Common questions and quick reference"),
]

def el(tag, **attrs):
    element = OxmlElement(tag)
    for key, val in attrs.items():
        element.set(qn("w:" + key), str(val))
    return element

def inline(paragraph, text):
    # Bold, italic and clickable HTTPS addresses. No external network requests.
    for part in re.split(r"(\*\*.*?\*\*|\*[^*]+\*|https://[^\s]+)", text):
        if not part:
            continue
        if part.startswith("https://"):
            href = part.rstrip(".,;")
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
            run = el("w:r")
            props = el("w:rPr")
            props.append(el("w:color", val="155A72"))
            props.append(el("w:u", val="single"))
            run.append(props)
            t = el("w:t")
            t.text = href
            run.append(t)
            link.append(run)
            paragraph._p.append(link)
            if len(href) < len(part):
                paragraph.add_run(part[len(href):])
        else:
            run = paragraph.add_run(part.strip("*") if part.startswith("*") else part)
            run.bold = part.startswith("**")
            run.italic = part.startswith("*") and not part.startswith("**")

def number_id(doc, start=1, bullet=False):
    root = doc.part.numbering_part.element
    abs_id = max([int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))] + [0]) + 1
    num_id = max([int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))] + [0]) + 1
    abstract = el("w:abstractNum", abstractNumId=abs_id)
    abstract.append(el("w:multiLevelType", val="singleLevel"))
    lvl = el("w:lvl", ilvl=0)
    lvl.append(el("w:start", val=start))
    lvl.append(el("w:numFmt", val="bullet" if bullet else "decimal"))
    lvl.append(el("w:lvlText", val="•" if bullet else "%1."))
    lvl.append(el("w:lvlJc", val="left"))
    lvl.append(el("w:suff", val="tab"))
    pp = el("w:pPr")
    tabs = el("w:tabs")
    tabs.append(el("w:tab", val="num", pos="490"))
    pp.append(tabs)
    pp.append(el("w:ind", left="490", hanging="490"))
    lvl.append(pp)
    rp = el("w:rPr")
    rp.append(el("w:rFonts", ascii=FONT, hAnsi=FONT))
    lvl.append(rp)
    abstract.append(lvl)
    root.append(abstract)
    num = el("w:num", numId=num_id)
    num.append(el("w:abstractNumId", val=abs_id))
    root.append(num)
    return num_id

def set_number(paragraph, num_id):
    pr = paragraph._p.get_or_add_pPr()
    num = el("w:numPr")
    num.append(el("w:ilvl", val="0"))
    num.append(el("w:numId", val=num_id))
    pr.append(num)

def style_doc(doc, short, index):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.67), Inches(.62)
    sec.left_margin = sec.right_margin = Inches(.7)
    sec.header_distance, sec.footer_distance = Inches(.25), Inches(.25)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "Header", "Footer", "List Paragraph"]:
        s = doc.styles[name]
        s.font.name = FONT
        s.font.color.rgb = RGBColor(0, 0, 0)
        rf = s.element.get_or_add_rPr()
        for child in list(rf):
            if child.tag == qn("w:color"):
                for a in ["themeColor", "themeTint", "themeShade"]:
                    child.attrib.pop(qn("w:" + a), None)
        rf.append(el("w:lang", val="en-NG"))
        s.paragraph_format.widow_control = True
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in [("Title", 25, 0, 10),("Heading 1", 20, 0, 10),("Heading 2", 13, 12, 6),("Heading 3", 11.5, 8, 5)]:
        s = doc.styles[name]
        s.font.size = Pt(size)
        s.font.bold = True
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.line_spacing = 1.02
        ppr = s.element.get_or_add_pPr()
        for b in list(ppr):
            if b.tag == qn("w:pBdr"):
                ppr.remove(b)
    doc.styles["Subtitle"].font.size = Pt(11.5)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(10)
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = True
    doc.styles["Caption"].paragraph_format.space_after = Pt(9)
    header = sec.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    header.paragraph_format.tab_stops.add_tab_stop(Inches(5.95))
    run = header.add_run(f"CESERVE User Guides\tGuide {index:02d}")
    run.font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.55))
    run = footer.add_run(f"{short}  ·  September 2026\t")
    run.font.size = Pt(8)
    field = el("w:fldSimple", instr="PAGE")
    r = el("w:r")
    rp = el("w:rPr"); rp.append(el("w:sz", val="16")); r.append(rp)
    t = el("w:t"); t.text = "1"; r.append(t); field.append(r)
    footer._p.append(field)
    doc.core_properties.author = "CESERVE"
    doc.core_properties.subject = "Courier Operating System end user instructions"
    doc.core_properties.keywords = "CESERVE, user guide, operations, configuration, booking"
    doc.core_properties.comments = "Prepared from the implemented application and repository contracts. Training examples only."
    doc.core_properties.created = datetime(2026, 9, 18, 0, 0, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 9, 18, 0, 0, tzinfo=timezone.utc)

def make_table(doc, rows):
    cols = len(rows[0])
    table = doc.add_table(rows=1, cols=cols)
    table.autofit = False
    widths = [2.10, 5.0] if cols == 2 else [.58, 5.87, .65]
    if cols == 2 and rows[0][0] in ["Page", "Task"]:
        widths = [.58, 6.52] if rows[0][0] == "Page" else [6.52, .58]
    if cols == 2 and rows[0][-1] == "Page":
        widths = [6.1, 1.0]
    for c, w in zip(table.columns, widths):
        c.width = Inches(w)
    tblPr = table._tbl.tblPr
    borders = el("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        borders.append(el("w:" + edge, val="single", sz="4", color="D9D9D9"))
    tblPr.append(borders)
    for i, rowvalues in enumerate(rows):
        row = table.rows[0] if i == 0 else table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(el("w:cantSplit"))
        if i == 0:
            trpr.append(el("w:tblHeader"))
        for j, val in enumerate(rowvalues):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margins = el("w:tcMar")
            for side in ["top", "bottom"]:
                margins.append(el("w:" + side, w="85", type="dxa"))
            for side in ["left", "right"]:
                margins.append(el("w:" + side, w="110", type="dxa"))
            cp.append(margins)
            if i == 0 or i % 2 == 0:
                cp.append(el("w:shd", fill="EDF0F2" if i == 0 else "FAFBFC", val="clear"))
            p = cell.paragraphs[0]
            if cols == 2 and rows[0][-1] == "Page" and j == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            inline(p, val)
            for r in p.runs:
                r.font.size = Pt(10)
                if i == 0: r.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    p.add_run().font.size = Pt(1)

def build(source, dest, short, index):
    doc = Document()
    style_doc(doc, short, index)
    lines = source.read_text().splitlines()
    i = 0
    first_body = True
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line == "<!-- page -->":
            # Put a page break on the next heading, avoiding a blank paragraph.
            i += 1
            while i < len(lines) and not lines[i].strip(): i += 1
            line = lines[i].strip()
            assert line.startswith("## "), line
            p = doc.add_paragraph(line[3:], "Heading 1")
            p.paragraph_format.page_break_before = True
            i += 1; continue
        if line.startswith("# "):
            title = line[2:]
            assert re.fullmatch(r"[A-Za-z0-9 ]+", title), title
            doc.add_paragraph(title, "Title")
            doc.core_properties.title = title
            i += 1; continue
        if line.startswith("### ") or line.startswith("## "):
            title = line.split(" ", 1)[1]
            assert re.fullmatch(r"[A-Za-z0-9 ]+", title), title
            doc.add_paragraph(title, "Heading 2" if line.startswith("### ") else "Heading 1")
            i += 1; continue
        if line.startswith("|"):
            rows=[]
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", x) for x in cells): rows.append(cells)
                i += 1
            make_table(doc, rows); continue
        if line.startswith("!["):
            m = re.fullmatch(r"!\[(.*?)\]\((.*?)\)\{width=([\d.]+)\}", line)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            shape = p.add_run().add_picture(str(source.parent / m[2]), width=Inches(float(m[3])))
            shape._inline.docPr.set("descr", m[1])
            i += 1
            if i < len(lines) and lines[i].strip().startswith("*"):
                caption=doc.add_paragraph(lines[i].strip().strip("*"), "Caption")
                caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
                i += 1
            continue
        if line.startswith("- ") or re.match(r"^\d+\. ", line):
            bullet = line.startswith("- ")
            start = 1 if bullet else int(line.split(".",1)[0])
            nid=number_id(doc,start,bullet)
            while i < len(lines):
                item=lines[i].strip()
                if not (item.startswith("- ") if bullet else re.match(r"^\d+\. ", item)): break
                text = item[2:] if bullet else item.split(". ",1)[1]
                p=doc.add_paragraph(style="List Paragraph")
                p.paragraph_format.space_after=Pt(6)
                p.paragraph_format.left_indent=Inches(.34)
                p.paragraph_format.first_line_indent=Inches(-.34)
                set_number(p,nid); inline(p,text)
                i+=1
            continue
        p=doc.add_paragraph(style="Subtitle" if first_body else "Normal")
        inline(p,line)
        first_body=False
        i+=1
    doc.save(dest)
    print(f"Built {dest.name}: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")

if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for index, (source, filename, short) in enumerate(GUIDES, 1):
        build(HERE / (source + ".md"), OUT / (filename + ".docx"), short, index)
